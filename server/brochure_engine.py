import os
import io
import time
import re
from typing import List, Optional
import fitz
import docx
from docx.shared import Pt, Twips, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

STANDARD_SIZES = [(595, 842), (612, 792), (842, 595), (792, 612)]  # A4, Letter (Portrait & Landscape)

FONT_MAP = {
    'rubik': 'Segoe UI',
    'poppin': 'Segoe UI',
    'montserrat': 'Segoe UI',
    'outfit': 'Segoe UI',
    'inter': 'Segoe UI',
    'roboto': 'Segoe UI',
    'open sans': 'Segoe UI',
    'lato': 'Segoe UI',
    'proxima nova': 'Segoe UI',
    'helvetica': 'Arial',
    'arialmt': 'Arial',
    'timesnewroman': 'Times New Roman',
    'times-roman': 'Times New Roman',
    'calibril': 'Calibri',
    'aptos': 'Aptos',
}

def clean_font_name(font_name: str, default_font: str = "Segoe UI") -> str:
    if not font_name:
        return default_font
    clean = re.sub(r'^[A-Z]{6}\+', '', font_name).strip()
    name_lower = clean.lower()
    for key, target in FONT_MAP.items():
        if key in name_lower:
            return target
    base = clean.split('-')[0].split(',')[0].strip()
    return base if base else default_font

def is_brochure_archetype(pdf_path: str, sample_pages: int = 4) -> bool:
    """
    Detector heurístico ultrarrápido (<50ms) de arquetipo Brochure / Folleto / Catálogo:
    Identifica documentos gráficos con dimensiones personalizadas, alta densidad
    vectorial (polígonos, banners, iconos) o dispersión multi-columna.
    """
    doc = None
    try:
        doc = fitz.open(pdf_path)
        score = 0
        pages_to_check = min(len(doc), sample_pages)
        if pages_to_check == 0:
            return False

        for i in range(pages_to_check):
            page = doc[i]
            w, h = round(page.rect.width), round(page.rect.height)
            
            # 1. Tamaño personalizado no estándar (Ej. 680.3 x 765.4 pt)
            is_standard = any(abs(w - sw) <= 12 and abs(h - sh) <= 12 for sw, sh in STANDARD_SIZES)
            if not is_standard:
                score += 2.5
                
            # 2. Densidad de elementos gráficos y vectoriales
            drawings = page.get_drawings()
            if len(drawings) >= 80:
                score += 3.0
            elif len(drawings) >= 30:
                score += 1.5

            images = page.get_images()
            if len(images) >= 3:
                score += 1.5
                
            # 3. Dispersión espacial de bloques de texto (múltiples ejes de columnas)
            blocks = page.get_text("blocks")
            text_blocks = [b for b in blocks if b[6] == 0]
            x0_coords = [round(b[0], -1) for b in text_blocks]
            if len(set(x0_coords)) >= 4:
                score += 2.0

        avg_score = score / pages_to_check
        return avg_score >= 3.0
    except Exception as e:
        print(f"Warning in is_brochure_archetype: {e}")
        return False
    finally:
        if doc:
            try:
                doc.close()
            except Exception:
                pass


class PDFBlackBrochureEngine:
    """
    Motor OpenXML de Posicionamiento Absoluto y Lienzo Gráfico Fiel
    Diseñado bajo arquitectura Gemini 4 Argon para replicar o superar a CloudConvert:
    - 0 tablas espurias (evita descalabro de columnas de revista).
    - 0 tabuladores '\\t' parásitos entre palabras.
    - Preservación milimétrica de dimensiones de página personalizadas.
    - Smart Background: Renderizado acelerado de vectores complejos en fondo de alta definición.
    - Párrafos posicionados con <w:framePr> nativo de Word (100% editables).
    """

    def __init__(self, pdf_path: str, output_docx: str, dpi: int = 150):
        self.pdf_path = pdf_path
        self.output_docx = output_docx
        self.dpi = dpi
        self.doc = fitz.open(pdf_path)
        self.word_doc = docx.Document()
        # Limpiar párrafo por defecto
        self.word_doc._body.clear_content()

    def _pt_to_twips(self, pt: float) -> int:
        return int(round(pt * 20))

    def _pt_to_emu(self, pt: float) -> int:
        return int(round(pt * 12700))

    def _render_clean_background(self, page_num: int) -> bytes:
        """
        Clona la página, oculta el texto y extrae el fondo vectorial e imágenes puras en PNG.
        """
        temp_pdf = fitz.open()
        temp_pdf.insert_pdf(self.doc, from_page=page_num, to_page=page_num)
        temp_page = temp_pdf[0]
        
        for block in temp_page.get_text("blocks"):
            if block[6] == 0:  # Tipo texto
                temp_page.add_redact_annot(block[:4])
        temp_page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
        
        pix = temp_page.get_pixmap(dpi=self.dpi, alpha=False)
        img_bytes = pix.tobytes("png")
        temp_pdf.close()
        return img_bytes

    def _add_background_drawingml(self, paragraph, img_bytes: bytes, width_pt: float, height_pt: float, page_idx: int):
        run = paragraph.add_run()
        inline_shape = run.add_picture(io.BytesIO(img_bytes))
        r_id = run._r.xpath('.//a:blip/@r:embed')[0]
        
        cx = self._pt_to_emu(width_pt)
        cy = self._pt_to_emu(height_pt)
        doc_id = 2000 + page_idx
        
        run._r.clear()
        
        anchor_xml = f"""
        <w:drawing {nsdecls('w', 'wp', 'a', 'pic', 'r')}>
            <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
                <wp:simplePos x="0" y="0"/>
                <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
                <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
                <wp:extent cx="{cx}" cy="{cy}"/>
                <wp:wrapNone/>
                <wp:docPr id="{doc_id}" name="PDFBlack_Background_{page_idx}"/>
                <a:graphic>
                    <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
                        <pic:pic>
                            <pic:nvPicPr>
                                <pic:cNvPr id="{doc_id}" name="BG_{page_idx}"/>
                                <pic:cNvPicPr/>
                            </pic:nvPicPr>
                            <pic:blipFill>
                                <a:blip r:embed="{r_id}"/>
                                <a:stretch><a:fillRect/></a:stretch>
                            </pic:blipFill>
                            <pic:spPr>
                                <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
                                <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                            </pic:spPr>
                        </pic:pic>
                    </a:graphicData>
                </a:graphic>
            </wp:anchor>
        </w:drawing>
        """
        run._r.append(parse_xml(anchor_xml))

    def _add_absolute_text_frame(self, block_dict: dict):
        bbox = block_dict["bbox"]
        x_twips = self._pt_to_twips(bbox[0])
        y_twips = self._pt_to_twips(bbox[1])
        w_twips = self._pt_to_twips(bbox[2] - bbox[0]) + 140  # Holgura de seguridad anti-wrap
        
        p = self.word_doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        
        # Espaciado cero para alineación milimétrica
        spacing = OxmlElement('w:spacing')
        spacing.set(qn('w:before'), '0')
        spacing.set(qn('w:after'), '0')
        spacing.set(qn('w:line'), '240')
        spacing.set(qn('w:lineRule'), 'auto')
        pPr.append(spacing)

        # Marco con coordenadas absolutas de página
        framePr = OxmlElement('w:framePr')
        framePr.set(qn('w:w'), str(w_twips))
        framePr.set(qn('w:hRule'), 'auto')
        framePr.set(qn('w:hSpace'), '0')
        framePr.set(qn('w:vSpace'), '0')
        framePr.set(qn('w:wrap'), 'notBeside')
        framePr.set(qn('w:vAnchor'), 'page')
        framePr.set(qn('w:hAnchor'), 'page')
        framePr.set(qn('w:x'), str(x_twips))
        framePr.set(qn('w:y'), str(y_twips))
        pPr.append(framePr)

        # Reconstrucción de líneas unificando fragmentos de la misma línea base horizontal
        prev_y0 = None
        for line in block_dict.get("lines", []):
            y0 = line["bbox"][1]
            if prev_y0 is not None and abs(y0 - prev_y0) > 3.5:
                p.add_run("\n")  # Salto de renglón visual real
            prev_y0 = y0
            
            for span in line.get("spans", []):
                clean_text = span.get("text", "").replace('\t', ' ').strip()
                if not clean_text:
                    continue
                run = p.add_run(clean_text + " ")
                font_name = clean_font_name(span.get("font", ""), "Segoe UI")
                run.font.name = font_name
                run.font.size = Pt(max(6.5, round(span.get("size", 10.0) * 2) / 2))
                
                c = span.get("color", 0)
                if c and c != 0:
                    r = (c >> 16) & 255
                    g = (c >> 8) & 255
                    b = c & 255
                    run.font.color.rgb = RGBColor(r, g, b)
                
                font_lower = span.get("font", "").lower()
                flags = span.get("flags", 0)
                if (flags & 2**4) or any(w in font_lower for w in ["bold", "black", "heavy", "medium"]):
                    run.bold = True
                if (flags & 2**1) or any(w in font_lower for w in ["italic", "oblique"]):
                    run.italic = True

    def convert(self, target_pages: Optional[List[int]] = None) -> str:
        start_time = time.time()
        pages_to_process = target_pages if target_pages is not None else list(range(len(self.doc)))
        total_pages = len(pages_to_process)
        
        for idx, page_num in enumerate(pages_to_process):
            if page_num < 0 or page_num >= len(self.doc):
                continue
            page = self.doc[page_num]
            w_pt, h_pt = page.rect.width, page.rect.height
            
            # 1. Fondo gráfico/vectorial limpio en PNG
            bg_p = self.word_doc.add_paragraph()
            bg_bytes = self._render_clean_background(page_num)
            self._add_background_drawingml(bg_p, bg_bytes, w_pt, h_pt, idx + 1)
            
            # 2. Bloques de texto posicionales
            text_page = page.get_text("dict", flags=fitz.TEXTFLAGS_TEXT)
            for block in text_page.get("blocks", []):
                if block.get("type") == 0:  # Bloque de texto
                    self._add_absolute_text_frame(block)
            
            # 3. Sección de página con geometría exacta
            if idx < total_pages - 1:
                section = self.word_doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
            else:
                section = self.word_doc.sections[-1]
                
            section.page_width = Twips(self._pt_to_twips(w_pt))
            section.page_height = Twips(self._pt_to_twips(h_pt))
            section.top_margin = Twips(0)
            section.bottom_margin = Twips(0)
            section.left_margin = Twips(0)
            section.right_margin = Twips(0)

        # Configurar la primera sección si existe
        if self.word_doc.sections:
            first_p = self.doc[pages_to_process[0]]
            first_sec = self.word_doc.sections[0]
            first_sec.page_width = Twips(self._pt_to_twips(first_p.rect.width))
            first_sec.page_height = Twips(self._pt_to_twips(first_p.rect.height))
            first_sec.top_margin = Twips(0)
            first_sec.bottom_margin = Twips(0)
            first_sec.left_margin = Twips(0)
            first_sec.right_margin = Twips(0)

        self.word_doc.save(self.output_docx)
        self.doc.close()
        elapsed = time.time() - start_time
        print(f"[BrochureEngine] Conversión exitosa de {total_pages} páginas en {elapsed:.2f}s.")
        return self.output_docx
